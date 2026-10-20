"""Local Excel -> semantic plan -> native BlackDoc prototype (Python 3.10+).

Excel-to-MD supplies the rules, facts and revision guardrails. This renderer
reads the workbook directly; Markdown is never an intermediate document.
"""
from __future__ import annotations

import argparse
import base64
from collections import Counter
from datetime import date, datetime, time
import hashlib
import importlib.util
import json
import mimetypes
from pathlib import Path
import posixpath
import re
import sys
from urllib.parse import quote
import uuid
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.cell.rich_text import CellRichText, TextBlock

VERSION = 1
NS = {
    "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "xdr": "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
}
KINDS = {"heading", "paragraph", "label", "compact_label", "ordered_item",
         "unordered_item", "condition", "continuation", "review", "skip"}


def digest(path):
    checksum = hashlib.sha256()
    with Path(path).open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            checksum.update(chunk)
    return checksum.hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8-sig"))


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # Exclusive creation also protects manually adjusted documents and plans.
    with path.open("x", encoding="utf-8", newline="\n") as stream:
        json.dump(value, stream, ensure_ascii=False, indent=2)
        stream.write("\n")


def load_converter(path):
    sys.dont_write_bytecode = True
    spec = importlib.util.spec_from_file_location("excel2md_source", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def relationships(archive, part):
    folder, name = posixpath.split(part)
    rel = posixpath.join(folder, "_rels", name + ".rels")
    if rel not in archive.namelist():
        return {}
    result = {}
    for entry in ET.fromstring(archive.read(rel)):
        if entry.get("TargetMode") == "External":
            continue
        target = entry.get("Target", "")
        result[entry.get("Id")] = (target.lstrip("/") if target.startswith("/")
                                   else posixpath.normpath(posixpath.join(folder, target)))
    return result


def drawing_manifest(workbook):
    """Read drawing relationships directly, including WPS images openpyxl skips."""
    images, unsupported = [], []
    with zipfile.ZipFile(workbook) as archive:
        book_rels = relationships(archive, "xl/workbook.xml")
        root = ET.fromstring(archive.read("xl/workbook.xml"))
        for sheet in root.findall("main:sheets/main:sheet", NS):
            name = sheet.get("name")
            part = book_rels[sheet.get("{" + NS["r"] + "}id")]
            sheet_root = ET.fromstring(archive.read(part))
            sheet_rels = relationships(archive, part)
            for drawing in sheet_root.findall("main:drawing", NS):
                drawing_part = sheet_rels[drawing.get("{" + NS["r"] + "}id")]
                drawing_rels = relationships(archive, drawing_part)
                for index, anchor in enumerate(ET.fromstring(archive.read(drawing_part))):
                    marker = anchor.find("xdr:from", NS)
                    pictures = anchor.findall("xdr:pic", NS)
                    if marker is None or not pictures:
                        unsupported.append({"sheet": name, "part": drawing_part,
                                            "anchor": index, "reason": "unsupported_drawing"})
                        continue
                    row = int(marker.findtext("xdr:row", namespaces=NS)) + 1
                    column = int(marker.findtext("xdr:col", namespaces=NS)) + 1
                    end = anchor.find("xdr:to", NS)
                    for picture in pictures:
                        blip = picture.find("xdr:blipFill/a:blip", NS)
                        media = drawing_rels.get(blip.get("{" + NS["r"] + "}embed")) if blip is not None else None
                        if media is None or media not in archive.namelist():
                            unsupported.append({"sheet": name, "row": row,
                                                "reason": "unresolved_picture"})
                            continue
                        images.append({"sheet": name, "row": row, "column": column,
                                       "end_row": int(end.findtext("xdr:row", namespaces=NS)) + 1 if end is not None else row,
                                       "media": media, "key": f"{drawing_part}:{index}:{media}"})
    return images, unsupported


def stable_id(*parts):
    return str(uuid.uuid5(uuid.NAMESPACE_URL, "blackdoc-excel:" + json.dumps(parts, ensure_ascii=False)))


def block(kind, key, content=None, props=None, children=None):
    return {"id": stable_id(*key), "type": kind, "props": props or {},
            **({"content": content} if content is not None else {}), "children": children or []}


def text(value, styles=None):
    return {"type": "text", "text": value, "styles": styles or {}}


def native_color(value, background=False):
    """Map Excel RGB to BlackDoc's theme-aware palette; retain the original in facts."""
    if not value:
        return "default"
    rgb = str(value).lstrip("#")[-6:]
    if not re.fullmatch("[0-9a-fA-F]{6}", rgb):
        return "default"
    r, g, b = [int(rgb[i:i + 2], 16) / 255 for i in (0, 2, 4)]
    high, low = max(r, g, b), min(r, g, b)
    if (not background and high < .18) or (background and low > .97):
        return "default"
    delta = high - low
    if delta < .08:
        return "gray"
    hue = ((g - b) / delta if high == r else (b - r) / delta + 2 if high == g else (r - g) / delta + 4) * 60 % 360
    return next(color for end, color in [(15, "red"), (45, "orange"), (75, "yellow"),
                (165, "green"), (255, "blue"), (300, "purple"), (345, "pink"), (361, "red")] if hue < end)


def displayed_value(cell, cached, reviews):
    value = cell.value
    if cell.data_type == "f":
        value = cached[cell.coordinate].value
        if value is None:
            reviews.append({"sheet": cell.parent.title, "cell": cell.coordinate,
                            "reason": "formula_without_cached_result", "formula": cell.value})
            return cell.value
    if value is None:
        return ""
    fmt = cell.number_format
    if isinstance(value, (datetime, date, time)):
        pattern = re.sub(r"\\(.)", r"\1", fmt).lower()
        pattern = pattern.replace('"', '')
        if re.fullmatch(r"[ymdhHsS0-9 /:.,-]+", pattern, re.IGNORECASE):
            hour_start = pattern.find('h')
            def token(match):
                raw = match.group()
                letter = raw[0]
                if letter == 'y':
                    return str(value.year) if len(raw) == 4 else f"{value.year % 100:02}"
                field = {'d': 'day', 'h': 'hour', 's': 'second',
                         'm': 'minute' if hour_start >= 0 and match.start() > hour_start else 'month'}[letter]
                number = getattr(value, field, 0)
                return str(number) if len(raw) == 1 else f"{number:02}"
            return re.sub(r"yyyy|yy|mm|m|dd|d|hh|h|ss|s", token, pattern)
        reviews.append({"sheet": cell.parent.title, "cell": cell.coordinate,
                        "reason": "unsupported_date_format", "format": fmt})
        return value.isoformat(sep=' ') if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)):
        simple = fmt.split(";")[0]
        if re.fullmatch(r"0+", simple):
            return f"{value:0{len(simple)}.0f}"
        match = re.fullmatch(r"(#,##)?0(?:\.(0+))?(%)?", simple)
        if match:
            places = len(match[2] or "")
            return format(value * (100 if match[3] else 1), f"{',' if match[1] else ''}.{places}f") + ("%" if match[3] else "")
        if fmt != "General":
            reviews.append({"sheet": cell.parent.title, "cell": cell.coordinate,
                            "reason": "unsupported_number_format", "format": fmt})
    return str(value)


class Renderer:
    def __init__(self, wb, cached, plan, images, archive, helper):
        self.wb, self.cached, self.plan = wb, cached, plan
        self.images, self.archive, self.helper = images, archive, helper
        self.source_key = plan["source_sha256"]
        self.mapping, self.skipped, self.reviews = {}, [], []
        self.links, self.image_keys, self.row_targets, self.sheet_targets = [], [], {}, {}
        self.id_keys = set()
        self.history_sheet = helper.find_revision_history_sheet(wb)
        self.history_layout = helper.find_history_layout(self.history_sheet) if self.history_sheet is not None else None

    def record(self, ws, cell, target):
        key = ws.title + "!" + cell.coordinate
        if key in self.mapping:
            raise ValueError("Source cell rendered twice: " + key)
        self.mapping[key] = target

    def styles(self, font, fill=None):
        styles = {key: True for key, enabled in [("bold", font.b), ("italic", font.i),
                  ("strike", font.strike), ("underline", font.u)] if enabled}
        color = native_color(self.helper.color_hex(font.color))
        if color != "default":
            styles["textColor"] = color
        if fill:
            styles["backgroundColor"] = native_color(fill, True)
        return styles

    def cell(self, ws, current, target, *, heading=False, label=False, omit_source_fill=False):
        self.record(ws, current, target)
        value = displayed_value(current, self.cached[ws.title], self.reviews)
        fill = None if heading or omit_source_fill else self.helper.fill_background_hex(current)
        content = []
        runs = current.value if isinstance(current.value, CellRichText) else [value]
        for run in runs:
            font = run.font if isinstance(run, TextBlock) else current.font
            styles = self.styles(font, fill)
            if heading:
                styles.pop("bold", None)
            elif label:
                styles["bold"] = True
            content.append(text(run.text if isinstance(run, TextBlock) else str(run), styles))
        if current.comment:
            self.reviews.append({"sheet": ws.title, "cell": current.coordinate,
                                 "reason": "cell_comment", "text": current.comment.text})
        link = current.hyperlink
        if link:
            result = {"type": "link", "href": link.target or "", "content": content}
            location = link.location or (link.target[1:] if link.target and link.target.startswith("#") else None)
            if location:
                self.links.append((result, ws.title, current.coordinate, location))
            return [result]
        return content

    def background(self, ws, row):
        ranges = self.helper.semantic_revision_blocks(self.plan, ws.title)
        revision = self.helper.revision_block_for_row(ranges, row)
        return native_color(self.helper.revision_background(revision), True) if revision else "default"

    def image(self, item, in_table=False, width=480):
        self.image_keys.append(item["key"])
        data = self.archive.read(item["media"])
        # WPS may store PNG bytes under a .jpeg media name.
        mime = ("image/png" if data.startswith(b"\x89PNG\r\n\x1a\n") else
                "image/jpeg" if data.startswith(b"\xff\xd8\xff") else mimetypes.guess_type(item["media"])[0])
        if mime not in {"image/png", "image/jpeg", "image/gif", "image/webp", "image/svg+xml", "image/bmp"}:
            raise ValueError("Unsupported embedded image: " + item["media"])
        url = f"data:{mime};base64," + base64.b64encode(data).decode("ascii")
        props = {"url": url, "name": Path(item["media"]).name, "previewWidth": width}
        if in_table:
            return {"type": "tableImage", "props": props}
        ws = self.wb[item["sheet"]]
        result = block("image", (self.source_key, item["key"]), props={**props,
                       "showPreview": True, "textAlignment": "left", "backgroundColor": self.background(ws, item["row"])})
        self.image_mapping.append({"sheet": item["sheet"], "row": item["row"],
                                   "column": item["column"], "block": result["id"]})
        return result

    def row(self, ws, row, decision):
        kind = decision.get("kind")
        if kind not in KINDS:
            raise ValueError(f"Invalid row kind: {ws.title}!{row}: {kind}")
        entries = [c for c in ws[row] if c.value is not None]
        if kind == "skip":
            if not decision.get("reason"):
                raise ValueError("Skipped rows need a reason")
            self.skipped.extend({"sheet": ws.title, "cell": c.coordinate, "reason": decision["reason"]} for c in entries)
            return None
        level = int(decision.get("level", 3))
        if kind == "heading" and not 3 <= level <= 6:
            raise ValueError("Worksheet headings must use levels 3 through 6")
        block_type = {"heading": "heading", "ordered_item": "numberedListItem",
                      "unordered_item": "bulletListItem", "condition": "bulletListItem"}.get(kind, "paragraph")
        result = block(block_type, (self.source_key, ws.title, "row", row), [],
                       {"backgroundColor": self.background(ws, row)})
        if kind == "heading":
            result["props"].update(level=level, isToggleable=True)
        if kind == "ordered_item" and entries and str(entries[0].value).isdigit():
            marker = entries.pop(0)
            result["props"]["start"] = int(marker.value)
            self.record(ws, marker, result["id"])
        for current in entries:
            if result["content"]:
                result["content"].append(text("　"))
            result["content"].extend(self.cell(ws, current, result["id"], heading=kind == "heading",
                                               label=kind in {"label", "compact_label"}))
        self.row_targets[(ws.title, row)] = result["id"]
        if kind == "review":
            self.reviews.append({"sheet": ws.title, "row": row, "reason": decision.get("reason", "semantic_review")})
        return result

    def table(self, ws, spec):
        start, end = int(spec["start_row"]), int(spec["end_row"])
        columns = spec["columns"]
        if not (1 <= start <= end <= ws.max_row) or not columns or columns != sorted(set(columns)) or any(not isinstance(c, int) or not 1 <= c <= ws.max_column for c in columns):
            raise ValueError("Invalid table range")
        header_rows = int(spec.get("header_rows", 1))
        if not 0 <= header_rows <= end - start + 1:
            raise ValueError("Invalid header_rows")
        result = block("table", (self.source_key, ws.title, "table", start),
                       {"type": "tableContent", "headerRows": header_rows, "rows": []})
        row_order = list(range(start, end + 1))
        if ws is self.history_sheet:
            layout = self.helper.find_history_layout(ws)
            if layout and start == layout[0] and header_rows and all(c in columns for c in layout[1:4]):
                # Reuse Excel-to-MD history semantics; move each complete source row.
                row_order[header_rows:] = sorted(
                    row_order[header_rows:],
                    key=lambda row: self.helper.parse_revision_date(str(ws.cell(row, layout[1]).value or "")),
                    reverse=True,
                )
        for row in row_order:
            history_header = (ws is self.history_sheet and self.history_layout
                              and start == self.history_layout[0] and row < start + header_rows)
            cells = []
            for column in columns:
                current = ws.cell(row, column)
                content = self.cell(ws, current, result["id"], omit_source_fill=history_header) if current.value is not None else []
                props = {"backgroundColor": self.background(ws, row), "textColor": "default",
                         "textAlignment": "left"}
                fill = None if history_header else self.helper.fill_background_hex(current)
                if props["backgroundColor"] == "default" and fill:
                    props["backgroundColor"] = native_color(fill, True)
                for image in self.images:
                    if (image["sheet"], image["row"], image["column"]) == (ws.title, row, column):
                        content.append(self.image(image, True, 240))
                        self.image_mapping.append({"sheet": ws.title, "row": row, "column": column, "block": result["id"]})
                cells.append({"type": "tableCell", "props": props, "content": content})
            # Outside cells are not silently appended to a guessed table column.
            extras = [c.coordinate for c in ws[row] if c.value is not None and c.column not in columns]
            if extras:
                raise ValueError(f"Table excludes source cells at {ws.title}!{extras}; include their columns")
            self.row_targets[(ws.title, row)] = result["id"]
            result["content"]["rows"].append({"cells": cells})
        return result

    def rows(self, ws, spec, start, end, exclude_images=False):
        result, parents = [], {}
        tables = {int(t["start_row"]): t for t in spec.get("tables", [])}
        row = start
        while row <= end:
            if row in tables:
                table = tables[row]
                if table["end_row"] > end:
                    raise ValueError("Table crosses a layout boundary")
                if not exclude_images:
                    result.extend(self.image(i) for i in self.images if i["sheet"] == ws.title
                                  and row <= i["row"] <= table["end_row"]
                                  and i["column"] not in table["columns"])
                result.append(self.table(ws, table))
                parents.clear()
                row = table["end_row"] + 1
                continue
            images = [i for i in self.images if i["sheet"] == ws.title and i["row"] == row]
            if not exclude_images:
                result.extend(self.image(i) for i in images)
            if any(c.value is not None for c in ws[row]):
                decision = spec.get("rows", {}).get(str(row))
                if not isinstance(decision, dict):
                    raise ValueError(f"Unplanned source row: {ws.title}!{row}")
                current = self.row(ws, row, decision)
                if current is not None:
                    depth = int(decision.get("indent", 0))
                    if depth < 0:
                        raise ValueError("Negative indentation")
                    for key in list(parents):
                        if key >= depth:
                            del parents[key]
                    if depth:
                        parent = parents.get(depth - 1)
                        if parent is None:
                            raise ValueError(f"Missing parent for indent {depth}: {ws.title}!{row}")
                        parent["children"].append(current)
                    else:
                        result.append(current)
                    if current["type"] in {"bulletListItem", "numberedListItem"}:
                        parents[depth] = current
            row += 1
        return result

    def resolve_links(self):
        for link, source_sheet, coordinate, location in self.links:
            target = None
            location = location.lstrip("#")
            match = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!\$?([A-Z]+)\$?(\d+)", location)
            if match:
                sheet = (match[1] or match[2]).replace("''", "'")
                cell_coord = match[3] + match[4]
                target = self.mapping.get(sheet + "!" + cell_coord)
                if not target and cell_coord == "A1":
                    target = self.sheet_targets.get(sheet)
            elif re.fullmatch(r"\$?[A-Z]+\$?\d+", location):
                target = self.mapping.get(source_sheet + "!" + location.replace("$", ""))
            else:
                # Excel/WPS revision links can target a text marker instead of a coordinate.
                marker = location.lstrip("_")
                matches = {self.mapping.get(ws.title + "!" + c.coordinate)
                           for ws in self.wb for row in ws for c in row
                           if c.value is not None and str(c.value).strip() == marker}
                matches.discard(None)
                if len(matches) == 1:
                    target = matches.pop()
            if target:
                link["href"] = "#block=" + quote(target)
            else:
                # Keep readable text in the document and the original target in the report.
                source = self.wb[source_sheet][coordinate]
                preserve = (self.history_sheet is self.wb[source_sheet] and self.history_layout
                            and source.row > self.history_layout[0]
                            and source.column in self.history_layout[4])
                if preserve:
                    link["href"] = source.hyperlink.target or "#" + source.hyperlink.location
                else:
                    link.pop("href", None)
                self.reviews.append({"sheet": source_sheet, "cell": coordinate,
                                     "reason": "unresolved_internal_link", "target": location})
        def flatten(value):
            if isinstance(value, list):
                output = []
                for item in value:
                    if isinstance(item, dict) and item.get("type") == "link" and "href" not in item:
                        output.extend(flatten(item["content"]))
                    else:
                        output.append(flatten(item))
                return output
            if isinstance(value, dict):
                return {k: flatten(v) for k, v in value.items()}
            return value
        return flatten(self.blocks)

    def render(self):
        self.image_mapping = []
        self.blocks = [block("heading", (self.source_key, "title"), [text(self.plan["title"])],
                             {"level": 1, "isToggleable": True})]
        worksheets = list(self.wb)
        if self.history_sheet is not None:
            worksheets.remove(self.history_sheet)
            worksheets.insert(0, self.history_sheet)
        for ws in worksheets:
            spec = self.plan.get("sheets", {}).get(ws.title)
            if not isinstance(spec, dict):
                raise ValueError("Missing worksheet plan: " + ws.title)
            previous_end = 0
            for table in sorted(spec.get("tables", []), key=lambda t: t["start_row"]):
                if table["start_row"] <= previous_end:
                    raise ValueError("Overlapping tables: " + ws.title)
                previous_end = table["end_row"]
            if spec.get("skip"):
                if not spec.get("reason") or any(i["sheet"] == ws.title for i in self.images):
                    raise ValueError("Skipped sheet needs a reason and cannot discard images")
                self.skipped.extend({"sheet": ws.title, "cell": c.coordinate, "reason": spec["reason"]}
                                    for row in ws for c in row if c.value is not None)
                continue
            heading = block("heading", (self.source_key, ws.title, "sheet"), [text(ws.title)],
                            {"level": 2, "isToggleable": True})
            self.blocks.append(heading)
            self.sheet_targets[ws.title] = heading["id"]
            max_row = max(ws.max_row, max((i["row"] for i in self.images if i["sheet"] == ws.title), default=0))
            position = 1
            for group in sorted(spec.get("image_groups", []), key=lambda g: g["start"]):
                start, end = group["start"], group["end"]
                separate = (group.get("render_as") == "separate"
                            or group.get("vertical_coverage") == "insufficient")
                if not position <= start <= end <= max_row or (not separate and group.get("layout") != "split"):
                    raise ValueError("Invalid or overlapping image group")
                self.blocks.extend(self.rows(ws, spec, position, start - 1))
                if separate:
                    self.blocks.extend(self.rows(ws, spec, start, end))
                    position = end + 1
                    continue
                pictures = [i for i in self.images if i["sheet"] == ws.title and start <= i["row"] <= end]
                if not pictures:
                    raise ValueError("Split layout requires pictures")
                left = [self.image(i) for i in pictures]
                right = self.rows(ws, spec, start, end, True)
                if not right:
                    raise ValueError("Split layout requires description blocks")
                self.blocks.append(block("splitPane", (self.source_key, ws.title, "split", start),
                    props={"leftWidth": 45, "rightHeight": 400}, children=[
                        block("splitColumn", (self.source_key, ws.title, "split-left", start), props={"side": "left"}, children=left),
                        block("splitColumn", (self.source_key, ws.title, "split-right", start), props={"side": "right"}, children=right)]))
                position = end + 1
            self.blocks.extend(self.rows(ws, spec, position, max_row))
            for review in spec.get("manual_review", []):
                self.reviews.append({"sheet": ws.title, **review})
        self.blocks = self.resolve_links()
        accounted = set(self.mapping) | {s["sheet"] + "!" + s["cell"] for s in self.skipped}
        expected = {ws.title + "!" + c.coordinate for ws in self.wb for row in ws for c in row if c.value is not None}
        if expected != accounted:
            raise ValueError("Source coverage mismatch: " + str(sorted(expected - accounted)[:10]))
        if Counter(self.image_keys) != Counter(i["key"] for i in self.images):
            raise ValueError("Missing or duplicate image placements")
        return self.blocks, {"source_cells": len(expected), "rendered_cells": len(self.mapping),
                             "skipped_cells": self.skipped, "images": len(self.image_keys),
                             "source_map": self.mapping, "image_map": self.image_mapping,
                             "manual_review": self.reviews}


def main():
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("mode", choices=["context", "render"])
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--converter", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    parser.add_argument("--reuse-context", type=Path, help="Enrich an existing fact extraction without repeating it")
    parser.add_argument("--plan", type=Path)
    args = parser.parse_args()
    if args.output.exists():
        parser.error("Output exists; choose a new staging path")
    workbook = args.workbook.resolve()
    source_hash = digest(workbook)
    converter_hash = digest(args.converter)
    rules_hash = digest(args.converter.parent / "excel2md.md")
    converter = load_converter(args.converter)
    images, unsupported = drawing_manifest(workbook)
    if args.mode == "context":
        if args.reuse_context:
            context = read_json(args.reuse_context)
            if context.get("source_sha256") != source_hash:
                parser.error("Reusable context must carry the matching workbook hash")
            if context.get("converter_sha256") != converter_hash:
                parser.error("Reusable context must come from the current source converter")
            context["rules"] = (args.converter.parent / "excel2md.md").read_text(encoding="utf-8")
        else:
            wb = openpyxl.load_workbook(workbook, data_only=False, rich_text=True)
            context = converter.build_semantic_context(wb, workbook)
        context["source_sha256"] = source_hash
        context["rules_sha256"] = rules_hash
        context["converter_sha256"] = converter_hash
        context["renderer_version"] = VERSION
        context["unsupported_drawings"] = unsupported
        for sheet in context["facts"]["sheets"]:
            sheet["images"] = [i for i in images if i["sheet"] == sheet["name"]]
        context["bdoc_contract"] = {
            "root_required": ["renderer_version", "source_sha256", "rules_sha256", "converter_sha256", "title", "sheets"],
            "rules": ["Use original text; every nonempty row needs a role or explicit table coverage.",
                      "Tables: {start_row, end_row, columns: [1-based source columns], header_rows: 0 or 1}.",
                      "No source cell outside the selected table columns may be silently dropped.",
                      "Rows: existing row roles, indent means nesting under a preceding list item.",
                      "Every skipped row or sheet needs a reason. Images cannot be skipped.",
                      "Optional image_groups: {start, end, layout: 'split'}; render_as: 'separate' or vertical_coverage: 'insufficient' forces source-order separate blocks.",
                      "Tables cannot cross image group boundaries or overlap other tables.",
                      "revision_blocks and heading guardrails follow the supplied excel2md rules."]}
        context["document_contract"] = {
            "version": 1, "plan_field": "document_version: 1",
            "sheet_fields": {"document": "ordered nodes", "omitted_cells": "array of {cell, reason}"},
            "node_kinds": ["heading", "paragraph", "bullet", "ordered", "code", "table", "image"],
            "content": "source refs: cell or {cell, separator: whitespace, bold, italic, format: iso_date}",
            "table": "{kind: table, rows: [[{content: refs, images: [{image: key, width}]}]], header_rows}",
            "rules": ["One source row can produce multiple nodes; cells may merge into one target cell.",
                      "Every nonempty source cell is rendered exactly once or omitted with a reason.",
                      "Every picture appears exactly once at its explicit semantic position.",
                      "Headings inside revision scopes, numeric UI labels and levels 5/6 need keep_as_heading.",
                      "Revision scope backgrounds override decorative source fills.",
                      "The same compiled tree feeds both adapters; adapters cannot infer structure."]}
        write_json(args.output, context)
        print(json.dumps({"context": str(args.output), "images": len(images), "unsupported_drawings": len(unsupported)}, ensure_ascii=False))
        return
    if args.plan is None:
        parser.error("render requires --plan")
    if args.output.suffix.lower() != ".bdoc":
        parser.error("render output must use the .bdoc extension")
    report_path = args.output.with_suffix(".report.json")
    if report_path.exists():
        parser.error("Report exists; choose a new staging path")
    plan = read_json(args.plan)
    shared_document = plan.get("document_version") is not None
    companion_paths = [args.output.with_suffix(ext) for ext in [".document.json", ".md"]] if shared_document else []
    if any(path.exists() for path in companion_paths):
        parser.error("Document tree or Markdown exists; choose a new staging path")
    if plan.get("source_sha256") != source_hash or plan.get("renderer_version") != VERSION:
        parser.error("Plan does not match the workbook or renderer version")
    if plan.get("rules_sha256") != rules_hash or plan.get("converter_sha256") != converter_hash:
        parser.error("Rules or source converter changed; review the plan against the new context")
    wb = openpyxl.load_workbook(workbook, data_only=False, rich_text=True)
    cached = openpyxl.load_workbook(workbook, data_only=True)
    validation = converter.validate_semantic_plan(wb, plan)
    if validation["status"] not in {"ok", "manual_review"}:
        raise ValueError("Invalid revision plan: " + json.dumps(validation, ensure_ascii=False))
    plan, heading_audit = converter.apply_heading_safety_policy(wb, plan)
    with zipfile.ZipFile(workbook) as archive:
        if shared_document:
            from excel_document import DocumentCompiler, bdoc, markdown
            tree, report = DocumentCompiler(wb, cached, plan, images, archive, converter).compile()
            blocks, markdown_output = bdoc(tree), markdown(tree)
        else:
            renderer = Renderer(wb, cached, plan, images, archive, converter)
            blocks, report = renderer.render()
    report.update(source_sha256=source_hash, renderer_version=VERSION,
                  rules_sha256=rules_hash, converter_sha256=converter_hash,
                  semantic_revision=validation, heading_audit=heading_audit,
                  unsupported_drawings=unsupported,
                  status="needs_review" if report["manual_review"] or unsupported else "generated")
    if digest(workbook) != source_hash:
        raise ValueError("Workbook changed during rendering; generate a new context")
    # Write only after source, image and revision checks have all succeeded.
    write_json(report_path, report)
    write_json(args.output, blocks)
    if shared_document:
        write_json(companion_paths[0], tree)
        with companion_paths[1].open("x", encoding="utf-8", newline="\n") as stream:
            stream.write(markdown_output)
    print(json.dumps({"output": str(args.output), "report": str(report_path),
                      "source_cells": report["source_cells"], "images": report["images"],
                      "manual_review": len(report["manual_review"]), "status": report["status"]}, ensure_ascii=False))


if __name__ == "__main__":
    main()
