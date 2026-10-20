from datetime import datetime
import json
from io import BytesIO
from pathlib import Path
import sys
from types import SimpleNamespace
import unittest
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from excel_document import DocumentCompiler, bdoc, markdown, plain, walk


def inputs():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Example'
    for coord, value in {'A1': 'Conditions', 'A2': 'Label', 'B2': 'Normal detail',
                         'A3': '2026.04.21 added', 'B3': 'Refresh', 'A4': 'Refresh detail',
                         'A5': 'Name', 'B5': 'Description', 'A6': 'Duration',
                         'B6': 'Seven days', 'D6': 'Configuration check',
                         'A7': 'Field one', 'A8': 'Field two', 'A9': 'Go to change'}.items():
        ws[coord] = value
    ws['A2'].font = Font(b=True)
    ws['A3'].font = Font(b=True, color='FFFF0000')
    ws['A5'].font = Font(b=True, color='FFFFFFFF')
    for coord in ['A3', 'B3', 'A4', 'A5']:
        ws[coord].fill = PatternFill('solid', fgColor='FF0080FF')
    ws['A9'].hyperlink = '#Example!A3'
    plan = {'source_sha256': 'synthetic-document', 'document_version': 1, 'title': 'Example document',
            'sheets': {'Example': {'revision_blocks': [{'start_row': 3, 'end_row': 4}], 'document': [
                {'key': 'title', 'kind': 'heading', 'level': 3, 'content': ['A1']},
                {'key': 'label', 'kind': 'paragraph', 'content': ['A2']},
                {'key': 'detail', 'kind': 'paragraph', 'content': ['B2']},
                {'key': 'marker', 'kind': 'paragraph', 'background': '#fff4cc', 'content': [{'cell': 'A3', 'bold': False}]},
                {'key': 'refresh', 'kind': 'bullet', 'background': '#fff4cc', 'content': ['B3'], 'children': [
                    {'key': 'detail-refresh', 'kind': 'paragraph', 'background': '#fff4cc', 'content': ['A4']}]},
                {'key': 'table', 'kind': 'table', 'rows': [[{'content': ['A5']}, {'content': ['B5']}],
                    [{'content': ['A6']}, {'content': ['B6', {'cell': 'D6', 'separator': '\n'}]}]]},
                {'key': 'fields', 'kind': 'code', 'content': ['A7', {'cell': 'A8', 'separator': '\n'}]},
                {'key': 'link', 'kind': 'paragraph', 'content': ['A9']},
            ]}}}
    helper = SimpleNamespace(
        find_revision_history_sheet=lambda book: None,
        color_hex=lambda color: '#' + color.rgb[-6:] if color and color.type == 'rgb' else None,
        fill_background_hex=lambda cell: '#' + cell.fill.fgColor.rgb[-6:] if cell.fill.fill_type else None,
        semantic_revision_blocks=lambda plan, sheet: plan['sheets'][sheet].get('revision_blocks', []),
    )
    return wb, plan, helper


def fixture():
    wb, plan, helper = inputs()
    return DocumentCompiler(wb, wb, plan, [], None, helper).compile()


class SharedDocumentTests(unittest.TestCase):
    def test_checked_in_adapter_fixtures_match_the_source_backed_tree(self):
        tree, _ = fixture()
        root = Path(__file__).parent / 'fixtures'
        self.assertEqual(json.loads((root / 'document.bdoc').read_text(encoding='utf-8')), bdoc(tree))
        self.assertEqual((root / 'document.md').read_text(encoding='utf-8'), markdown(tree))

    def test_guarded_heading_requires_explicit_opt_in(self):
        wb, plan, helper = inputs()
        title = plan['sheets']['Example']['document'][0]
        title['level'] = 5
        with self.assertRaisesRegex(ValueError, 'Guarded heading'):
            DocumentCompiler(wb, wb, plan, [], None, helper).compile()
        title['keep_as_heading'] = True
        tree, _ = DocumentCompiler(wb, wb, plan, [], None, helper).compile()
        self.assertEqual(tree['nodes'][2]['level'], 5)

    def test_one_tree_preserves_split_merge_styles_code_links_and_revision_structure(self):
        tree, report = fixture()
        native = bdoc(tree)
        md = markdown(tree)
        nodes = list(walk(tree['nodes']))
        self.assertEqual(report['source_cells'], report['rendered_cells'])
        self.assertEqual([b['props']['level'] for b in native if b['type'] == 'heading'], [1, 2, 3])
        self.assertEqual(next(n for n in nodes if plain(n.get('content', [])) == 'Normal detail')['content'][0]['styles'], {})
        marker = next(n for n in nodes if plain(n.get('content', [])) == '2026.04.21 added')
        self.assertEqual(marker['content'][0]['styles'], {'textColor': '#FF0000'})
        bullet = next(b for b in native if b['type'] == 'bulletListItem')
        self.assertEqual(bullet['children'][0]['content'][0]['text'], 'Refresh detail')
        table = next(b for b in native if b['type'] == 'table')
        self.assertEqual(plain(table['content']['rows'][1]['cells'][1]['content']), 'Seven days\nConfiguration check')
        self.assertEqual(table['content']['rows'][0]['cells'][0]['content'][0]['styles']['textColor'], '#FFFFFF')
        self.assertEqual(next(b for b in native if b['type'] == 'codeBlock')['content'][0]['text'], 'Field one\nField two')
        self.assertIn('- <span style="display:block;background:#fff4cc">Refresh</span>', md)
        self.assertIn('```text\nField one\nField two\n```', md)
        self.assertIn('href="#' + marker['id'] + '"', md)
        self.assertEqual(native[-1]['content'][0]['href'], '#block=' + marker['id'])

    def test_duplicate_missing_source_and_unhighlighted_revision_are_rejected(self):
        for alteration, error in [
            ('duplicate', 'rendered twice'), ('missing', 'coverage mismatch'),
            ('revision', 'missing semantic background'), ('key', 'Duplicate document key')]:
            with self.subTest(alteration=alteration):
                wb, plan, helper = inputs()
                nodes = plan['sheets']['Example']['document']
                if alteration == 'duplicate':
                    nodes.append({'key': 'duplicate', 'kind': 'paragraph', 'content': ['A1']})
                elif alteration == 'missing':
                    nodes.pop()
                elif alteration == 'revision':
                    del nodes[3]['background']
                else:
                    nodes[1]['key'] = nodes[0]['key']
                with self.assertRaisesRegex(ValueError, error):
                    DocumentCompiler(wb, wb, plan, [], None, helper).compile()

    def test_image_can_belong_to_table_cell_and_move_with_its_semantic_parent(self):
        wb, plan, helper = inputs()
        pictures = [{'sheet': 'Example', 'row': 99, 'column': 1, 'key': 'picture', 'media': 'xl/media/pic.png'}]
        node = plan['sheets']['Example']['document'][5]
        node['rows'][1][0]['images'] = [{'image': 'picture', 'width': 70}]
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('xl/media/pic.png', b'synthetic-image-bytes')
        with zipfile.ZipFile(buffer) as archive:
            tree, report = DocumentCompiler(wb, wb, plan, pictures, archive, helper).compile()
        self.assertEqual(report['images'], 1)
        table = next(b for b in bdoc(tree) if b['type'] == 'table')
        self.assertEqual(table['content']['rows'][1]['cells'][0]['content'][-1]['type'], 'tableImage')
        self.assertIn('width:70px', markdown(tree))

    def test_date_display_is_decided_once_for_both_adapters(self):
        wb, plan, helper = inputs()
        wb.active['B2'] = datetime(2026, 3, 1)
        wb.active['B2'].number_format = 'mm/dd'
        plan['sheets']['Example']['document'][2]['content'] = [{'cell': 'B2', 'format': 'iso_date'}]
        tree, _ = DocumentCompiler(wb, wb, plan, [], None, helper).compile()
        self.assertIn('2026-03-01', markdown(tree))
        self.assertTrue(any(b.get('content') == [{'type': 'text', 'text': '2026-03-01', 'styles': {}}] for b in bdoc(tree)))


if __name__ == '__main__':
    unittest.main()
