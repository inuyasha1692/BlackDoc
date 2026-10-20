import base64
from datetime import datetime
import importlib.util
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
import unittest
import uuid
import zipfile

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.styles import PatternFill

spec = importlib.util.spec_from_file_location('excel_to_bdoc', Path(__file__).parents[1] / 'excel_to_bdoc.py')
native = importlib.util.module_from_spec(spec)
spec.loader.exec_module(native)

# A data-only substitute for the unchanged Excel-to-MD helper interface.
helper = SimpleNamespace(
    find_revision_history_sheet=lambda wb: None,
    color_hex=lambda c: '#' + c.rgb[-6:] if c and c.type == 'rgb' else None,
    fill_background_hex=lambda c: '#' + c.fill.fgColor.rgb[-6:] if c.fill.fill_type and c.fill.fgColor.type == 'rgb' else None,
    semantic_revision_blocks=lambda plan, sheet: plan['sheets'][sheet].get('revision_blocks', []),
    revision_block_for_row=lambda ranges, row: next((r for r in ranges if r['start_row'] <= row <= r['end_row']), None),
    revision_background=lambda r: r.get('background', '#fff4cc') if r else None,
)
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+jRZkAAAAASUVORK5CYII=')


def fixture_inputs():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Interface'
    for coordinate, value in {
        'A1': 'Overview', 'A2': 'Parent item', 'A3': 'Nested detail',
        'A4': 'Revision marker', 'A5': CellRichText(TextBlock(InlineFont(b=True, color='FFFF0000'), 'Release '), 'detail'),
        'A6': 'Name', 'B6': 'Preview', 'A7': 'Main', 'B7': 'Help',
        'A9': 'Image description', 'A10': 'Back to overview',
    }.items():
        ws[coordinate] = value
    ws['A10'].hyperlink = "#'Interface'!A1"
    rows = {str(r): {'kind': k} for r, k in [(1, 'heading'), (2, 'unordered_item'),
            (3, 'continuation'), (4, 'label'), (5, 'paragraph'), (9, 'paragraph'), (10, 'paragraph')]}
    rows['1']['level'] = 3
    rows['3']['indent'] = 1
    plan = {'source_sha256': 'synthetic-fixture', 'title': 'Synthetic Excel', 'sheets': {'Interface': {
        'rows': rows, 'tables': [{'start_row': 6, 'end_row': 7, 'columns': [1, 2]}],
        'image_groups': [{'start': 9, 'end': 10, 'layout': 'split'}],
        'revision_blocks': [{'start_row': 4, 'end_row': 5, 'background': '#fff4cc'}],
    }}}
    images = [{'sheet': 'Interface', 'row': row, 'column': col,
               'key': str(row), 'media': 'xl/media/pic.png'} for row, col in [(7, 2), (9, 1)]]
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, 'w') as archive:
        archive.writestr('xl/media/pic.png', PNG)
    return wb, plan, images, buffer


def make_fixture():
    wb, plan, images, buffer = fixture_inputs()
    with zipfile.ZipFile(buffer) as archive:
        return native.Renderer(wb, wb, plan, images, archive, helper).render()


class NativeExcelTests(unittest.TestCase):
    def test_picture_mime_uses_png_bytes_even_with_jpeg_extension(self):
        wb, plan, images, buffer = fixture_inputs()
        for item in images:
            item['media'] = 'xl/media/pic.jpeg'
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('xl/media/pic.jpeg', PNG)
        with zipfile.ZipFile(buffer) as archive:
            blocks, _ = native.Renderer(wb, wb, plan, images, archive, helper).render()
        split = next(b for b in blocks if b['type'] == 'splitPane')
        self.assertTrue(split['children'][0]['children'][0]['props']['url'].startswith('data:image/png;base64,'))

    def test_explicit_separation_overrides_split_and_preserves_all_content(self):
        for decision in [{'render_as': 'separate'}, {'vertical_coverage': 'insufficient'}]:
            with self.subTest(decision=decision):
                wb, plan, images, buffer = fixture_inputs()
                plan['sheets']['Interface']['image_groups'][0].update(decision)
                with zipfile.ZipFile(buffer) as archive:
                    blocks, report = native.Renderer(wb, wb, plan, images, archive, helper).render()
                self.assertFalse(any(b['type'] == 'splitPane' for b in blocks))
                self.assertEqual(report['images'], 2)
                self.assertEqual(report['source_cells'], report['rendered_cells'])
                picture_index = next(i for i, b in enumerate(blocks) if b['type'] == 'image')
                self.assertEqual(blocks[picture_index + 1]['content'][0]['text'], 'Image description')

    def test_history_moves_first_and_only_extra_column_keeps_unresolved_link(self):
        wb = Workbook()
        wb.active.title = 'Body'
        wb.active['A1'] = 'Body text'
        history = wb.create_sheet('History')
        history.append(['Date', 'Author', 'Change', 'Ticket'])
        history.append(['2026.01.02', 'A', 'Missing section', 'Original ticket'])
        history['C1'] = CellRichText(TextBlock(InlineFont(b=True), 'Change'))
        for coordinate in ['A1', 'B1', 'C1', 'A2']:
            history[coordinate].fill = PatternFill('solid', fgColor='FF0080FF')
        history['C2'].hyperlink = '#Missing!A2'
        history['D2'].hyperlink = '#Missing!A3'
        plan = {'source_sha256': 'history-links', 'title': 'History links', 'sheets': {
            'Body': {'rows': {'1': {'kind': 'paragraph'}}},
            'History': {'rows': {}, 'tables': [{'start_row': 1, 'end_row': 2, 'columns': [1, 2, 3, 4]}]},
        }}
        history_helper = SimpleNamespace(**{**vars(helper),
            'find_revision_history_sheet': lambda workbook: workbook['History'],
            'find_history_layout': lambda sheet: (1, 1, 2, 3, [4]),
            'parse_revision_date': lambda value: (2026, 1, 2),
        })
        blocks, report = native.Renderer(wb, wb, plan, [], None, history_helper).render()
        self.assertEqual(blocks[1]['content'][0]['text'], 'History')
        header = blocks[2]['content']['rows'][0]['cells']
        self.assertTrue(all(cell['props']['backgroundColor'] == 'default' for cell in header))
        self.assertTrue(all('backgroundColor' not in item['styles'] for cell in header for item in cell['content']))
        self.assertTrue(header[2]['content'][0]['styles']['bold'])
        cells = blocks[2]['content']['rows'][1]['cells']
        self.assertEqual(cells[0]['props']['backgroundColor'], 'blue')
        self.assertEqual(cells[2]['content'], [native.text('Missing section')])
        self.assertEqual(cells[3]['content'][0]['href'], '#Missing!A3')
        self.assertEqual(len(report['manual_review']), 2)

    def test_headings_drop_source_fill_but_keep_links_and_revision_scope(self):
        wb, plan, images, buffer = fixture_inputs()
        ws = wb['Interface']
        ws['A1'] = CellRichText(TextBlock(InlineFont(b=True, color='FFFF0000'), 'Overview'))
        ws['A1'].hyperlink = 'https://example.com'
        for coordinate in ['A1', 'A2']:
            ws[coordinate].fill = PatternFill('solid', fgColor='FF00FF00')
        plan['sheets']['Interface']['revision_blocks'].append(
            {'start_row': 1, 'end_row': 1, 'background': '#fff4cc'})
        with zipfile.ZipFile(buffer) as archive:
            blocks, _ = native.Renderer(wb, wb, plan, images, archive, helper).render()
        heading = next(b for b in blocks if b['type'] == 'heading' and b['props'].get('level') == 3)
        self.assertEqual(heading['props']['backgroundColor'], 'yellow')
        link = heading['content'][0]
        self.assertEqual(link['href'], 'https://example.com')
        self.assertEqual(link['content'][0]['styles'], {'textColor': 'red'})
        paragraph = next(b for b in blocks if b['type'] == 'bulletListItem')
        self.assertEqual(paragraph['content'][0]['styles']['backgroundColor'], 'green')

    def test_history_sorts_complete_rows_stably_and_leaves_other_tables_in_source_order(self):
        wb = Workbook()
        ws = wb.active
        ws.title = 'History'
        for row in [('Updated', 'Author', 'Change'), ('2025.12.31', 'A', 'Old'),
                    ('2026.01.02', 'B', 'New'), ('2026.01.02', 'C', 'Same date'),
                    ('Undated', 'D', 'Unknown')]:
            ws.append(row)
        ws['C3'].hyperlink = 'https://example.com/new'
        ordinary = wb.copy_worksheet(ws)
        ordinary.title = 'Ordinary'
        plan = {'source_sha256': 'history-fixture', 'title': 'History test', 'sheets': {
            sheet.title: {'rows': {}, 'tables': [{'start_row': 1, 'end_row': 5, 'columns': [1, 2, 3]}]}
            for sheet in wb}}
        history_helper = SimpleNamespace(**{**vars(helper),
            'find_revision_history_sheet': lambda workbook: workbook['History'],
            'find_history_layout': lambda sheet: (1, 1, 2, 3, []),
            'parse_revision_date': lambda value: {'2025.12.31': (2025, 12, 31),
                '2026.01.02': (2026, 1, 2)}.get(value, (0, 0, 0)),
        })
        buffer = BytesIO()
        with zipfile.ZipFile(buffer, 'w') as archive:
            archive.writestr('xl/media/pic.png', PNG)
        images = [{'sheet': 'History', 'row': 3, 'column': 3, 'key': 'new', 'media': 'xl/media/pic.png'}]
        with zipfile.ZipFile(buffer) as archive:
            blocks, report = native.Renderer(wb, wb, plan, images, archive, history_helper).render()
        history, other = [b['content']['rows'] for b in blocks if b['type'] == 'table']
        self.assertEqual([r['cells'][1]['content'][0]['text'] for r in history],
                         ['Author', 'B', 'C', 'A', 'D'])
        self.assertEqual([r['cells'][1]['content'][0]['text'] for r in other],
                         ['Author', 'A', 'B', 'C', 'D'])
        new_content = history[1]['cells'][2]['content']
        self.assertEqual(new_content[0]['content'][0]['text'], 'New')
        self.assertEqual(new_content[0]['href'], 'https://example.com/new')
        self.assertEqual(new_content[1]['type'], 'tableImage')
        self.assertEqual(report['rendered_cells'], report['source_cells'])

    def test_native_structure_images_styles_links_and_source_coverage(self):
        blocks, report = make_fixture()
        self.assertEqual(report['source_cells'], report['rendered_cells'])
        self.assertEqual(report['images'], 2)
        self.assertEqual(report['manual_review'], [])
        parent = next(b for b in blocks if b['type'] == 'bulletListItem')
        self.assertEqual(parent['children'][0]['content'][0]['text'], 'Nested detail')
        revision = next(b for b in blocks if b.get('content', [{}])[0:1] == [native.text('Revision marker', {'bold': True})])
        self.assertEqual(revision['props']['backgroundColor'], 'yellow')
        table = next(b for b in blocks if b['type'] == 'table')
        self.assertEqual(table['content']['rows'][1]['cells'][1]['content'][-1]['type'], 'tableImage')
        split = next(b for b in blocks if b['type'] == 'splitPane')
        self.assertTrue(split['children'][0]['children'][0]['props']['url'].startswith('data:image/png;base64,'))
        self.assertEqual(split['children'][1]['children'][1]['content'][0]['href'], '#block=' + report['source_map']['Interface!A1'])

    def test_unplanned_nonempty_row_is_rejected(self):
        wb, plan, images, buffer = fixture_inputs()
        del plan['sheets']['Interface']['rows']['3']
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, 'Unplanned source row'):
            native.Renderer(wb, wb, plan, images, archive, helper).render()

    def test_excluded_table_cell_is_rejected(self):
        wb, plan, images, buffer = fixture_inputs()
        plan['sheets']['Interface']['tables'][0]['columns'] = [1]
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, 'excludes source cells'):
            native.Renderer(wb, wb, plan, images, archive, helper).render()

    def test_unplaced_picture_is_rejected(self):
        wb, plan, images, buffer = fixture_inputs()
        images.append({'sheet': 'Missing', 'row': 1, 'column': 1, 'key': 'unplaced', 'media': 'xl/media/pic.png'})
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, 'Missing or duplicate image'):
            native.Renderer(wb, wb, plan, images, archive, helper).render()

    def test_existing_output_is_protected(self):
        output = Path(__file__).resolve().parents[2] / 'output' / 'excel-tests'
        output.mkdir(parents=True, exist_ok=True)
        target = output / (str(uuid.uuid4()) + '.bdoc')
        try:
            native.write_json(target, ['manual reference'])
            with self.assertRaises(FileExistsError):
                native.write_json(target, ['replacement'])
            self.assertEqual(native.read_json(target), ['manual reference'])
        finally:
            target.unlink(missing_ok=True)

    def test_overlapping_tables_and_unjustified_skip_are_rejected(self):
        wb, plan, images, buffer = fixture_inputs()
        plan['sheets']['Interface']['tables'].append({'start_row': 7, 'end_row': 8, 'columns': [1]})
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, 'Overlapping tables'):
            native.Renderer(wb, wb, plan, images, archive, helper).render()
        plan['sheets']['Interface']['tables'].pop()
        plan['sheets']['Interface']['rows']['1'] = {'kind': 'skip'}
        with zipfile.ZipFile(buffer) as archive, self.assertRaisesRegex(ValueError, 'need a reason'):
            native.Renderer(wb, wb, plan, images, archive, helper).render()

    def test_unresolved_link_preserves_text_and_reports_original_target(self):
        wb, plan, images, buffer = fixture_inputs()
        wb['Interface']['A10'].hyperlink = "#'Missing'!A1"
        with zipfile.ZipFile(buffer) as archive:
            blocks, report = native.Renderer(wb, wb, plan, images, archive, helper).render()
        content = next(b for b in blocks if b['type'] == 'splitPane')['children'][1]['children'][1]['content']
        self.assertEqual(content, [native.text('Back to overview')])
        self.assertEqual(report['manual_review'][0]['target'], "'Missing'!A1")

    def test_displayed_dates_percentages_leading_zeros_and_missing_formula_cache(self):
        ws = Workbook().active
        reviews = []
        ws['A1'] = datetime(2026, 3, 1, 4, 5, 6)
        ws['A1'].number_format = 'mm/dd'
        self.assertEqual(native.displayed_value(ws['A1'], ws, reviews), '03/01')
        ws['A1'].number_format = r'yyyy/mm/dd\ hh:mm:ss'
        self.assertEqual(native.displayed_value(ws['A1'], ws, reviews), '2026/03/01 04:05:06')
        ws['A2'] = .125
        ws['A2'].number_format = '0.0%'
        self.assertEqual(native.displayed_value(ws['A2'], ws, reviews), '12.5%')
        ws['A3'] = 7
        ws['A3'].number_format = '0000'
        self.assertEqual(native.displayed_value(ws['A3'], ws, reviews), '0007')
        ws['A4'] = '=SUM(A2:A3)'
        self.assertEqual(native.displayed_value(ws['A4'], Workbook().active, reviews), '=SUM(A2:A3)')
        self.assertEqual(reviews[-1]['reason'], 'formula_without_cached_result')


if __name__ == '__main__':
    unittest.main()
